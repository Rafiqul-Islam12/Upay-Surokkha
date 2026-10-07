"""STEP 8: impact model (ScamShield ROI, Agent Copilot KPIs, Cost Advisor A/B pilot design).

Run from anywhere:   python src/impact_model.py [--no-xlsx] [--out-dir results]

Inputs : config/impact_assumptions.json   (EDIT THIS FILE or the blue cells in the XLSX)
Outputs: results/step8_impact_model.xlsx  (live formulas, editable blue cells)
         results/step8_assumptions.csv, step8_scenario_results.csv, step8_sensitivity_tornado.csv,
         step8_sensitivity_grid.csv, step8_agent_copilot_kpis.csv, step8_cost_advisor_ab_design.csv,
         step8_sources.csv, step8_results.md

EVERY input is either an ASSUMPTION (invented placeholder), a CONVENTION (statistical default) or
MEASURED_ON_SYNTHETIC (our own Step 2/3 result on synthetic data). None of them is real upay data.

Core formula (ScamShield, per month):
    prevented loss = transactions x scam rate x average loss x recall x intervention success
    net benefit    = prevented loss - false-alarm cost - analyst time cost - customer-friction cost - run cost
Set intervention success to 100% to obtain the literal "transactions x scam rate x loss x recall" figure
(it is also shown as its own row).
"""
import argparse
import csv
import json
import math
import sys
from pathlib import Path
from statistics import NormalDist

ROOT = Path(__file__).resolve().parent.parent
CONFIG = ROOT / "config" / "impact_assumptions.json"
SCEN = ["bad", "medium", "good"]
MODEL_GROUPS_EXCLUDED = set()   # all numeric assumptions feed the ScamShield model


# ----------------------------------------------------------------------------------------------
# 1. Assumptions
# ----------------------------------------------------------------------------------------------
def load_config(path=CONFIG):
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def scenario_inputs(cfg, scen):
    """{assumption id: value} for one scenario."""
    return {a["id"]: a[scen] for a in cfg["assumptions"]}


# ----------------------------------------------------------------------------------------------
# 2. ScamShield model (pure Python; the XLSX re-implements the same formulas)
# ----------------------------------------------------------------------------------------------
def compute(p):
    """Monthly impact for one input dict p (ids V1..H1). Returns a dict of results."""
    N, s, L, r, e = p["V1"], p["R1"], p["R2"], p["R3"], p["R4"]
    far, share, minutes = p["F1"], p["A1"], p["A2"]
    scams = N * s
    legit = N - scams
    gross_loss = scams * L
    literal = gross_loss * r                      # volume x rate x loss x recall
    prevented = literal * e
    tp = scams * r
    fn = scams - tp
    fp = legit * far
    flags = tp + fp
    precision = tp / flags if flags else 0.0
    reviews = flags * share
    cost_per_min = p["A3"] / (p["A4"] * 60 * p["A5"])
    cost_per_review = cost_per_min * minutes
    analyst_cost = reviews * cost_per_review
    fa_cost = fp * (p["F2"] + p["F3"] * p["F4"])
    friction = fp * (p["C1"] * p["C2"] + p["C3"] * p["C4"])
    run = p["K2"]
    net = prevented - fa_cost - analyst_cost - friction - run
    one_time = p["K1"]
    horizon = p["H1"]
    payback = one_time / net if net > 0 else None
    roi = (net * horizon - one_time) / one_time if one_time else None
    fte = reviews * minutes / (p["A4"] * 60 * p["A5"])
    # all-in cost of ONE false alarm (incl. the analyst review it triggers)
    cfa = p["F2"] + p["F3"] * p["F4"] + p["C1"] * p["C2"] + p["C3"] * p["C4"] + share * cost_per_review
    gain_per_tp = L * e - share * cost_per_review          # value of one correctly flagged scam, net of its review
    num = N * s * r * gain_per_tp - run
    be_far = num / (N * (1 - s) * cfa) if num > 0 and cfa > 0 else None
    den = N * (r * gain_per_tp + far * cfa)
    be_rate = (N * far * cfa + run) / den if den > 0 else None
    return {
        "scams": scams, "legit": legit, "gross_loss": gross_loss, "literal_prevented": literal,
        "prevented_loss": prevented, "tp": tp, "fn": fn, "fp": fp, "flags": flags, "precision": precision,
        "reviews": reviews, "cost_per_min": cost_per_min, "cost_per_review": cost_per_review,
        "analyst_cost": analyst_cost, "false_alarm_cost": fa_cost, "friction_cost": friction,
        "run_cost": run, "net_benefit": net, "one_time": one_time, "payback_months": payback,
        "roi": roi, "analyst_fte": fte, "cost_per_false_alarm_all_in": cfa,
        "breakeven_far": be_far, "breakeven_scam_rate": be_rate,
        "share_loss_prevented": prevented / gross_loss if gross_loss else 0.0,
        "cumulative_net_over_horizon": net * horizon - one_time,
    }


def net_closed_form(p, recall=None, far=None, scam_rate=None, loss=None):
    """Net benefit with optional overrides; used by sensitivity grids and tests."""
    q = dict(p)
    if recall is not None: q["R3"] = recall
    if far is not None: q["F1"] = far
    if scam_rate is not None: q["R1"] = scam_rate
    if loss is not None: q["R2"] = loss
    return compute(q)["net_benefit"]


def tornado(cfg, scen="medium", delta=0.25):
    """One-at-a-time +/-delta change of every input; effect on monthly net benefit."""
    base_p = scenario_inputs(cfg, scen)
    base = compute(base_p)["net_benefit"]
    rows = []
    for a in cfg["assumptions"]:
        if a["id"] in ("H1", "A4"):
            continue
        out = {}
        for tag, f in (("down", 1 - delta), ("up", 1 + delta)):
            q = dict(base_p)
            q[a["id"]] = base_p[a["id"]] * f
            if a["unit"] == "fraction":
                q[a["id"]] = min(q[a["id"]], 1.0)
            out[tag] = compute(q)["net_benefit"]
        rows.append({"id": a["id"], "assumption": a["name"], "base_value": base_p[a["id"]],
                     "net_at_minus_25pct": round(out["down"]), "net_at_plus_25pct": round(out["up"]),
                     "swing_bdt": round(abs(out["up"] - out["down"])), "base_net": round(base)})
    return sorted(rows, key=lambda r: -r["swing_bdt"])


RECALL_AXIS = [0.10, 0.18, 0.30, 0.447, 0.60, 0.75, 0.859]
FAR_AXIS = [0.005, 0.010, 0.0195, 0.0224, 0.040, 0.0677]
RATE_AXIS = [0.0001, 0.0002, 0.0003, 0.0006, 0.0010]
LOSS_AXIS = [2000, 4000, 6000, 10000, 15000]


# ----------------------------------------------------------------------------------------------
# 3. Agent Copilot KPIs and Cost Advisor A/B design (definitions + calculators)
# ----------------------------------------------------------------------------------------------
KPIS = [
    dict(kpi="Stock-out hours", unit="hours / agent / month", better="lower",
         definition="Hours in which the agent could not serve a normal cash-out because own cash or e-float was below the minimum service amount, or a cash-out request was declined for insufficient agent balance.",
         formula="stock-out hours = number of clock hours with (balance < min service amount) OR (>= 1 declined-for-liquidity request); averaged per agent per month",
         data_needed="Hourly agent cash/e-float snapshots; log of cash-out requests declined for liquidity; agent opening hours",
         baseline_example=24.0, pilot_example=16.0, target_example=-0.20,
         guardrail="Do not count hours when the agent was closed. Count only declines caused by liquidity, not by network or fraud holds."),
    dict(kpi="Completed transactions", unit="cash-out transactions / agent / week", better="higher",
         definition="Cash-out transactions completed at the agent, plus the liquidity-failure rate (share of cash-out attempts that failed because the agent lacked cash/float).",
         formula="completed = count of successful cash-outs per agent-week; liquidity-failure rate = failed-for-liquidity / (completed + failed-for-liquidity)",
         data_needed="Transaction log with status and failure reason code, per agent",
         baseline_example=310.0, pilot_example=325.0, target_example=0.03,
         guardrail="Compare same weeks of the month (salary week vs non-salary week) between arms."),
    dict(kpi="Idle cash", unit="BDT / agent (average daily excess buffer)", better="lower",
         definition="Average end-of-day cash + float held above what the next 7 days of observed demand needed (excess buffer). Lower is better only if stock-out hours do not rise.",
         formula="idle cash = mean over days of max(0, end-of-day balance - peak observed 7-day-ahead demand)",
         data_needed="End-of-day balances; realised daily cash-out demand per agent",
         baseline_example=60000.0, pilot_example=52000.0, target_example=-0.10,
         guardrail="Report idle cash and stock-out hours together; a drop in idle cash with more stock-outs is a failure, not a saving."),
    dict(kpi="Forecast error vs 7-day-average baseline", unit="MAE, BDT per agent-day", better="lower",
         definition="Mean absolute error of the 7-day forecast against realised cash-out demand, compared with the simple 7-day average.",
         formula="MAE = mean(|forecast - actual|); improvement = 1 - MAE_model / MAE_baseline",
         data_needed="Daily forecasts stored at issue time; realised daily demand",
         baseline_example=12256.0, pilot_example=10662.0, target_example=-0.10,
         guardrail="Example values are our SYNTHETIC Step-1 numbers (about 13% better), not real agent data."),
    dict(kpi="Warning usefulness", unit="share of HIGH warnings followed by a real shortfall", better="higher",
         definition="Precision of HIGH cash-shortage warnings and the share of warned agents who actually rebalanced before the predicted date.",
         formula="warning precision = HIGH warnings followed by a stock-out or near stock-out within 7 days / all HIGH warnings",
         data_needed="Warning log; rebalancing (cash-in / float purchase) log; stock-out log",
         baseline_example=0.30, pilot_example=0.45, target_example=0.10,
         guardrail="Too many unnecessary warnings make agents ignore the tool; track warnings per agent per week."),
]

AB_DESIGN = [
    ("Question", "How much do customers really save in fees, and how many small cash-outs really move to digital payments, when they see a Cost Advisor message? No migration rate is assumed: it is an OUTPUT of the pilot."),
    ("Unit of randomisation", "Customer, stratified by Cost Advisor segment (low / regular / frequent) and by area (urban / rural). Use the agent or village as the unit instead if customers share agents and spill-over is likely."),
    ("Eligibility", "Customers with at least 2 cash-outs under the small-cash-out limit in the 8-12 week pre-period, active in the last 30 days, not flagged by ScamShield, with message consent."),
    ("Arms", "Control = no message (true holdout). Treatment = Cost Advisor message. Optional arm B = message plus a small incentive, only if the team wants to price the incentive."),
    ("Allocation", "1:1 random assignment with a fixed seed written down before launch; check balance on pre-period fees, segment and area before launch (A/A check)."),
    ("Duration", "At least 8 weeks (two or more salary-week cycles). Do not stop early because an interim result looks good."),
    ("Primary metric", "Net fees paid on cash-outs per customer-month (BDT), treatment minus control, intention-to-treat. Variance reduction with the pre-period value (CUPED) if the correlation is known."),
    ("Secondary metrics", "Migration = share of eligible small cash-out value that moved to digital payments (measured, not assumed); count and value of digital payments; cash-out count."),
    ("Upay-side metric", "Net revenue per customer-month: fees on cash-outs + merchant discount earned - incentives. A customer saving is a revenue transfer unless merchant revenue makes up for it."),
    ("Guardrails", "Complaint rate, message opt-outs, total transaction value (must not fall), active-customer rate, cash-out failure rate. Any guardrail breach pauses the arm."),
    ("Analysis", "Pre-registered. Difference in means with 95% confidence interval (template in the workbook). Report ITT first; treatment-on-treated only as a supplement."),
    ("Decision rule", "Scale up only if the LOWER bound of the 95% CI of the saving is above the minimum worth acting on AND no guardrail is breached. Scale-up forecasts use the observed migration, never a fixed one."),
    ("Ethics and wording", "Message states a possible saving honestly and is optional advice; no pressure, no hiding of cash-out. Control customers get the message after the pilot if it works."),
]


def sample_size(sigma, mde, alpha, power, rho, attrition):
    """Customers per arm and in total for a two-sample difference in means (normal approximation)."""
    nd = NormalDist()
    z = nd.inv_cdf(1 - alpha / 2) + nd.inv_cdf(power)
    n_arm = 2 * z * z * sigma * sigma * (1 - rho * rho) / (mde * mde)
    n_arm_c = math.ceil(n_arm)
    total = math.ceil(2 * n_arm_c / (1 - attrition))
    return n_arm_c, total


def ab_params(cfg):
    return {b["id"]: b["value"] for b in cfg["ab_pilot"]}


# ----------------------------------------------------------------------------------------------
# 4. CSV export
# ----------------------------------------------------------------------------------------------
RESULT_ROWS = [
    ("scams", "Scam transactions per month", "count"), ("legit", "Legitimate transactions per month", "count"),
    ("gross_loss", "Total scam loss before ScamShield", "BDT"),
    ("literal_prevented", "Literal formula: transactions x rate x loss x recall", "BDT"),
    ("prevented_loss", "Prevented loss (x intervention success)", "BDT"),
    ("tp", "Scams flagged (true positives)", "count"), ("fn", "Scams missed", "count"),
    ("fp", "False alarms", "count"), ("flags", "Total flags", "count"), ("precision", "Precision at real prevalence", "fraction"),
    ("reviews", "Cases reviewed by analysts", "count"), ("analyst_fte", "Analysts needed", "FTE"),
    ("false_alarm_cost", "False-alarm cost (notification + support)", "BDT"),
    ("analyst_cost", "Analyst time cost", "BDT"), ("friction_cost", "Customer-friction cost", "BDT"),
    ("run_cost", "Run cost", "BDT"), ("net_benefit", "Net monthly benefit", "BDT"),
    ("one_time", "One-time cost", "BDT"), ("payback_months", "Payback", "months"),
    ("roi", "ROI over horizon", "fraction"), ("cumulative_net_over_horizon", "Cumulative net over horizon after one-time cost", "BDT"),
    ("cost_per_false_alarm_all_in", "All-in cost of one false alarm", "BDT"),
    ("breakeven_far", "Break-even false-alarm rate", "fraction"), ("breakeven_scam_rate", "Break-even scam rate", "fraction"),
]


def write_csv(path, header, rows):
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(rows)


def source_text(cfg, ids):
    return "; ".join(f"{i}: {cfg['sources'][i]['title']}" for i in ids)


def export_csvs(cfg, out):
    res = {s: compute(scenario_inputs(cfg, s)) for s in SCEN}
    write_csv(out / "step8_assumptions.csv",
              ["id", "group", "assumption", "unit", "bad", "medium", "good", "status", "basis", "public_sources_to_cite"],
              [[a["id"], a["group"], a["name"], a["unit"], a["bad"], a["medium"], a["good"], a["kind"], a["basis"],
                source_text(cfg, a["sources"])] for a in cfg["assumptions"]]
              + [[b["id"], "A/B pilot (Cost Advisor)", b["name"], b["unit"], b["value"], b["value"], b["value"], b["kind"],
                  b["basis"], source_text(cfg, b["sources"])] for b in cfg["ab_pilot"]])
    write_csv(out / "step8_scenario_results.csv", ["metric", "unit", "bad", "medium", "good"],
              [[lab, unit] + [("" if res[s][k] is None else round(res[s][k], 6)) for s in SCEN] for k, lab, unit in RESULT_ROWS])
    t = tornado(cfg)
    write_csv(out / "step8_sensitivity_tornado.csv", list(t[0].keys()), [list(r.values()) for r in t])
    base = scenario_inputs(cfg, "medium")
    rows = [["recall x false-alarm rate (medium)", f"recall={rc}", f"far={fa}", round(net_closed_form(base, recall=rc, far=fa))]
            for fa in FAR_AXIS for rc in RECALL_AXIS]
    rows += [["scam rate x average loss (medium)", f"scam_rate={sr}", f"avg_loss={lo}", round(net_closed_form(base, scam_rate=sr, loss=lo))]
             for lo in LOSS_AXIS for sr in RATE_AXIS]
    write_csv(out / "step8_sensitivity_grid.csv", ["grid", "x", "y", "net_monthly_benefit_bdt"], rows)
    write_csv(out / "step8_agent_copilot_kpis.csv",
              ["kpi", "unit", "better", "definition", "formula", "data_needed", "guardrail", "baseline_EXAMPLE", "pilot_EXAMPLE", "target_ASSUMPTION"],
              [[k["kpi"], k["unit"], k["better"], k["definition"], k["formula"], k["data_needed"], k["guardrail"],
                k["baseline_example"], k["pilot_example"], k["target_example"]] for k in KPIS])
    ab = ab_params(cfg)
    n_arm, total = sample_size(ab["B3"], ab["B4"], ab["B1"], ab["B2"], ab["B5"], ab["B6"])
    write_csv(out / "step8_cost_advisor_ab_design.csv", ["element", "design"],
              [list(r) for r in AB_DESIGN] + [["Sample size per arm", n_arm], ["Total customers to enrol", total]])
    used = {}
    for a in cfg["assumptions"] + cfg["ab_pilot"]:
        for sid in a["sources"]:
            used.setdefault(sid, []).append(a["id"])
    write_csv(out / "step8_sources.csv", ["id", "source", "where", "what_to_take_from_it", "used_for_inputs", "note"],
              [[sid, s["title"], s["url"], s["use"], ", ".join(used.get(sid, [])), source_note(sid)]
               for sid, s in cfg["sources"].items()])
    return res


def source_note(sid):
    if sid == "S1":
        return "Link appeared in a web search on 2026-10-07 with data up to May 2026. Open it and confirm the table before citing."
    if sid in ("S7", "S9"):
        return "Internal source."
    return "Link/edition NOT re-verified. Open the source, note the exact figure, edition and date, then cite."


def write_markdown(cfg, res, out):
    def f(x, kind):
        if x is None: return "n/a"
        if kind == "pct": return f"{x*100:.2f}%"
        if kind == "pct3": return f"{x*100:.4f}%"
        if kind == "m": return f"{x/1e6:,.2f} M"
        if kind == "n": return f"{x:,.1f}"
        return f"{x:,.0f}"
    lines = ["# Step 8: impact model (ScamShield ROI, Agent Copilot KPIs, Cost Advisor A/B pilot)", "",
             "**Every number below comes from assumptions, not from real upay data.** Inputs marked ASSUMPTION are invented placeholders;",
             "recall and false-alarm rate are our own synthetic-data results (Step 2 / Step 3). Edit `config/impact_assumptions.json`",
             "(or the blue cells in `results/step8_impact_model.xlsx`) and rerun `python src/impact_model.py`.", "",
             "## ScamShield: monthly impact (BDT million unless stated)", "",
             "| Metric | Bad | Medium | Good |", "|---|---|---|---|"]
    def row(lab, k, kind):
        lines.append(f"| {lab} | " + " | ".join(f(res[s][k], kind) for s in SCEN) + " |")
    row("Prevented loss", "prevented_loss", "m"); row("False-alarm cost", "false_alarm_cost", "m")
    row("Analyst time cost", "analyst_cost", "m"); row("Customer-friction cost", "friction_cost", "m")
    row("Run cost", "run_cost", "m"); row("**Net monthly benefit**", "net_benefit", "m")
    row("False alarms per month (count)", "fp", "n"); row("Precision at real prevalence", "precision", "pct")
    row("Analysts needed (FTE)", "analyst_fte", "n"); row("Payback (months)", "payback_months", "n")
    row("ROI over horizon", "roi", "pct"); row("Break-even false-alarm rate", "breakeven_far", "pct")
    row("Break-even scam rate", "breakeven_scam_rate", "pct3")
    lines += ["", "`n/a` = no payback (monthly net benefit is zero or negative) or no break-even exists. ROI below -100% means the system loses more over the horizon than the whole investment.", "",
              "## Reading it honestly (generated from the numbers above)", ""]
    for sc in SCEN:
        r = res[sc]
        verdict = "earns more than it costs" if r["net_benefit"] > 0 else "costs more than it earns"
        pb = "never pays back" if r["payback_months"] is None else f"pays back in {r['payback_months']:.1f} months"
        be = "no break-even exists" if r["breakeven_far"] is None else f"break-even false-alarm rate {r['breakeven_far']*100:.2f}% vs assumed {cfg_val(cfg,'F1',sc)*100:.2f}%"
        lines.append(f"- **{sc.capitalize()}**: net {f(r['net_benefit'],'m')} per month ({verdict}), {pb}; {be}; "
                     f"needs about {r['analyst_fte']:,.0f} analysts if every flag is reviewed.")
    lines += ["- Recall and false-alarm rate in these scenarios are our own synthetic-data results: good = Step 2 test split, medium = adapted-scam row, bad = external validation (Step 3). They do not predict real performance.",
              "- At a realistic (very low) scam rate almost every flag is a false alarm (see the precision row), so the false-alarm rate and the analyst cost per case dominate the result. Reviewing every flag by hand is the biggest cost line; triage and a lower false-alarm rate are the levers (see `step8_sensitivity_tornado.csv` and the Sensitivity sheet).",
              "- Step 2 used a flat 20 BDT per false alarm. Here it is decomposed; the all-in value per scenario is the row 'All-in cost of one false alarm' in `step8_scenario_results.csv`.", "",
              "## Agent Copilot", "",
              "KPIs, formulas, data needed and guardrails: `step8_agent_copilot_kpis.csv` and the `AgentCopilot_KPIs` sheet. Example baseline / pilot values are placeholders (the forecast-error row uses our synthetic Step 1 numbers). The dataset has no cash balances, so stock-out hours and idle cash cannot be measured on synthetic data.", "",
              "## Cost Advisor A/B pilot", ""]
    ab = ab_params(cfg)
    n_arm, total = sample_size(ab["B3"], ab["B4"], ab["B1"], ab["B2"], ab["B5"], ab["B6"])
    lines += [f"Design: `step8_cost_advisor_ab_design.csv`. With the placeholder inputs (sigma {ab['B3']} BDT, minimum detectable saving {ab['B4']} BDT per customer-month) the pilot needs about {n_arm:,} customers per arm ({total:,} to enrol). No migration rate is assumed; it is measured.", "",
              "## Which public sources to cite for each input", "", "See `step8_assumptions.csv` (column `public_sources_to_cite`) and `step8_sources.csv`.",
              "Only S1 was seen live; every other link must be opened and the exact figure and edition recorded before citing."]
    (out / "step8_results.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


def cfg_val(cfg, aid, scen):
    return next(a[scen] for a in cfg["assumptions"] if a["id"] == aid)


# ----------------------------------------------------------------------------------------------
# 5. XLSX (live formulas)
# ----------------------------------------------------------------------------------------------
def build_xlsx(cfg, path):
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter as L_

    FN = "Arial"
    f_norm = Font(name=FN, size=10)
    f_bold = Font(name=FN, size=10, bold=True)
    f_head = Font(name=FN, size=10, bold=True, color="FFFFFF")
    f_title = Font(name=FN, size=14, bold=True)
    f_in = Font(name=FN, size=10, color="0000FF")          # inputs
    f_link = Font(name=FN, size=10, color="008000")        # cross-sheet links
    f_warn = Font(name=FN, size=10, bold=True, color="C00000")
    fill_head = PatternFill("solid", fgColor="0B3F8F")
    fill_in = PatternFill("solid", fgColor="FFF2CC")
    fill_key = PatternFill("solid", fgColor="E2EFDA")
    fill_sec = PatternFill("solid", fgColor="D9E1F2")
    thin = Side(style="thin", color="BFBFBF")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")

    wb = Workbook()

    def setw(ws, widths):
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[L_(i)].width = w

    def head(ws, row, labels):
        for i, t in enumerate(labels, 1):
            c = ws.cell(row=row, column=i, value=t)
            c.font, c.fill, c.border, c.alignment = f_head, fill_head, box, Alignment(wrap_text=True, vertical="center")

    def put(ws, ref, value, font=f_norm, fmt=None, fill=None, border=True, align=None):
        c = ws[ref]
        c.value = value
        c.font = font
        if fmt: c.number_format = fmt
        if fill: c.fill = fill
        if border: c.border = box
        if align: c.alignment = align
        return c

    # ---------------- Read_Me ----------------
    ws = wb.active
    ws.title = "Read_Me"
    setw(ws, [26, 110])
    ws["A1"].value, ws["A1"].font = "upay Surokkha+ : Step 8 impact model", f_title
    lines = [
        ("Purpose", "Estimates the monthly money impact and payback of ScamShield; defines KPIs for Agent Copilot; designs an A/B pilot for Cost Advisor."),
        ("WARNING", "NO number here is real upay data. Inputs are ASSUMPTIONS (invented placeholders), CONVENTIONS (statistical defaults) or MEASURED_ON_SYNTHETIC (our own Step 2/3 results on synthetic data)."),
        ("Blue text, yellow fill", "Editable input. Change only these cells (sheets Assumptions, CostAdvisor_ABPilot, AgentCopilot_KPIs, and the axis cells in Sensitivity)."),
        ("Black text", "Formula. Do not overwrite."),
        ("Green text", "Link to another sheet."),
        ("Scenarios", "Bad / Medium / Good. Bad = pessimistic for benefit (low scam rate, low recall, high false-alarm rate, high costs). Good = optimistic."),
        ("Core formula", "Prevented loss = monthly transactions x scam rate x average loss x recall x intervention success. Row 'Literal formula' on ScamShield_Model shows it without the intervention factor (set intervention success to 100% to match it)."),
        ("Net benefit", "Prevented loss - false-alarm cost - analyst time cost - customer-friction cost - monthly run cost."),
        ("Payback / ROI", "Payback = one-time cost / monthly net benefit (only if net > 0). ROI = (net x horizon - one-time cost) / one-time cost."),
        ("Sheets", "Assumptions | ScamShield_Model | Sensitivity | AgentCopilot_KPIs | CostAdvisor_ABPilot | Sources"),
        ("Rebuild", "python src/impact_model.py   (rewrites this workbook and the CSV files from config/impact_assumptions.json)"),
        ("Currency", "BDT (Bangladeshi taka). Fractions are stored as fractions (0.02 = 2%)."),
    ]
    for i, (a, b) in enumerate(lines, 3):
        put(ws, f"A{i}", a, f_warn if a == "WARNING" else f_bold, align=wrap)
        put(ws, f"B{i}", b, f_warn if a == "WARNING" else f_norm, align=wrap)
        ws.row_dimensions[i].height = 32
    put(ws, "A16", "Colour legend", f_bold)
    put(ws, "B16", "Example input cell", f_in, fill=fill_in)
    put(ws, "B17", "Example formula cell", f_norm)
    put(ws, "B18", "Example cross-sheet link", f_link)

    # ---------------- Assumptions ----------------
    wa = wb.create_sheet("Assumptions")
    setw(wa, [6, 18, 52, 20, 14, 14, 14, 24, 62, 60])
    wa["A1"].value, wa["A1"].font = "Assumptions (edit the blue cells)", f_title
    wa["A2"].value, wa["A2"].font = ("Status column: ASSUMPTION = invented placeholder; MEASURED_ON_SYNTHETIC = our own result on synthetic data, not real-world. "
                                    "Sources column = public sources to cite when you replace the placeholder with a real figure."), f_warn
    head(wa, 4, ["ID", "Group", "Assumption", "Unit", "Bad", "Medium", "Good", "Status", "Basis / how it was set", "Public sources to cite (see Sources sheet)"])
    arow = {}
    r = 5
    for a in cfg["assumptions"]:
        arow[a["id"]] = r
        fmt = "0.00%" if a["unit"] == "fraction" else ("#,##0.00" if a["unit"] in ("BDT", "minutes") and isinstance(a["bad"], float) and a["bad"] < 100 else "#,##0")
        if a["id"] in ("R1", "C3"):
            fmt = "0.000%"
        put(wa, f"A{r}", a["id"], f_bold)
        put(wa, f"B{r}", a["group"], align=wrap)
        put(wa, f"C{r}", a["name"], align=wrap)
        put(wa, f"D{r}", a["unit"], align=wrap)
        for col, sc in zip("EFG", SCEN):
            put(wa, f"{col}{r}", a[sc], f_in, fmt, fill_in)
        put(wa, f"H{r}", a["kind"], f_bold if a["kind"] != "ASSUMPTION" else f_norm)
        put(wa, f"I{r}", a["basis"], align=wrap)
        put(wa, f"J{r}", source_text(cfg, a["sources"]) or "(team decision)", align=wrap)
        wa.row_dimensions[r].height = 54
        r += 1
    wa.freeze_panes = "E5"
    ACOL = {"bad": "E", "medium": "F", "good": "G"}

    # ---------------- ScamShield_Model ----------------
    wm = wb.create_sheet("ScamShield_Model")
    setw(wm, [58, 16, 18, 18, 18, 70])
    wm["A1"].value, wm["A1"].font = "ScamShield impact model (monthly, BDT)", f_title
    wm["A2"].value, wm["A2"].font = "All cells are formulas driven by the Assumptions sheet. Inputs are NOT real upay data.", f_warn
    head(wm, 4, ["Line", "Unit", "Bad", "Medium", "Good", "Formula / note"])
    MC = {"bad": "C", "medium": "D", "good": "E"}
    mrow = {}
    cur = [5]

    def section(title):
        for c in "ABCDEF":
            wm[f"{c}{cur[0]}"].fill = fill_sec
            wm[f"{c}{cur[0]}"].border = box
        wm[f"A{cur[0]}"].value, wm[f"A{cur[0]}"].font = title, f_bold
        cur[0] += 1

    def line(key, label, unit, tmpl, fmt="#,##0", note="", link=None, key_row=False):
        """tmpl uses {AID} for assumption refs (link rows) and [key] for model-row refs."""
        row = cur[0]
        mrow[key] = row
        put(wm, f"A{row}", label, f_bold if key_row else f_norm, align=wrap)
        put(wm, f"B{row}", unit)
        for sc in SCEN:
            col = MC[sc]
            if link:
                formula = f"=Assumptions!{ACOL[sc]}{arow[link]}"
                font = f_link
            else:
                formula = "=" + tmpl
                for k in sorted(mrow, key=len, reverse=True):
                    formula = formula.replace(f"[{k}]", f"{col}{mrow[k]}")
                font = f_bold if key_row else f_norm
            put(wm, f"{col}{row}", formula, font, fmt, fill_key if key_row else None)
        put(wm, f"F{row}", note, align=wrap)
        cur[0] += 1

    section("A. Inputs used (linked from Assumptions)")
    line("N", "Monthly transactions scored", "tx / month", "", "#,##0", "V1", link="V1")
    line("s", "Scam rate", "fraction", "", "0.000%", "R1", link="R1")
    line("L", "Average loss per scam", "BDT", "", "#,##0", "R2", link="R2")
    line("r", "Recall", "fraction", "", "0.0%", "R3 (own Step 2/3 result on synthetic data)", link="R3")
    line("e", "Intervention success", "fraction", "", "0.0%", "R4", link="R4")
    line("far", "False-alarm rate", "fraction", "", "0.00%", "F1 (own Step 2/3 result on synthetic data)", link="F1")
    line("share", "Share of flags reviewed by analyst", "fraction", "", "0%", "A1", link="A1")
    line("mins", "Review minutes per case", "minutes", "", "#,##0.0", "A2", link="A2")
    line("sal", "Analyst monthly cost", "BDT / month", "", "#,##0", "A3", link="A3")
    line("hrs", "Paid hours per month", "hours", "", "#,##0", "A4", link="A4")
    line("util", "Utilization", "fraction", "", "0%", "A5", link="A5")
    line("notif", "Notification cost per flag", "BDT", "", "#,##0.00", "F2", link="F2")
    line("scr", "Support-contact rate per false alarm", "fraction", "", "0.0%", "F3", link="F3")
    line("scc", "Cost per support contact", "BDT", "", "#,##0.00", "F4", link="F4")
    line("aband", "Abandonment rate after false alarm", "fraction", "", "0.0%", "C1", link="C1")
    line("rev", "Revenue lost per abandoned tx", "BDT", "", "#,##0.00", "C2", link="C2")
    line("churn", "Churn probability per false alarm", "fraction", "", "0.000%", "C3", link="C3")
    line("cval", "Annual margin per customer", "BDT / year", "", "#,##0", "C4", link="C4")
    line("capex", "One-time cost", "BDT", "", "#,##0", "K1", link="K1")
    line("opex", "Monthly run cost", "BDT / month", "", "#,##0", "K2", link="K2")
    line("hor", "ROI horizon", "months", "", "#,##0", "H1", link="H1")

    section("B. Prevented loss")
    line("scams", "Scam transactions per month", "count", "[N]*[s]", "#,##0", "volume x scam rate")
    line("legit", "Legitimate transactions per month", "count", "[N]-[scams]", "#,##0")
    line("gross", "Total scam loss without ScamShield", "BDT", "[scams]*[L]", "#,##0", "scams x average loss")
    line("literal", "Literal formula: volume x rate x loss x recall", "BDT", "[gross]*[r]", "#,##0", "the formula as specified (intervention success = 100%)")
    line("prev", "Prevented loss", "BDT", "[literal]*[e]", "#,##0", "literal x intervention success", key_row=True)
    line("tp", "Scams flagged (true positives)", "count", "[scams]*[r]", "#,##0")
    line("fn", "Scams missed", "count", "[scams]-[tp]", "#,##0")

    section("C. Costs")
    line("fp", "False alarms per month", "count", "[legit]*[far]", "#,##0", "legitimate transactions x false-alarm rate")
    line("flags", "Total flags", "count", "[tp]+[fp]", "#,##0")
    line("prec", "Precision at this prevalence", "fraction", "IF([flags]=0,0,[tp]/[flags])", "0.0%", "share of flags that are real scams")
    line("fa", "False-alarm cost (notification + support contact)", "BDT", "[fp]*([notif]+[scr]*[scc])", "#,##0", "false alarms x (notification + contact rate x contact cost)")
    line("reviews", "Cases reviewed by analysts", "count", "[flags]*[share]", "#,##0", "true and false flags both need review")
    line("cpm", "Analyst cost per productive minute", "BDT / min", "[sal]/([hrs]*60*[util])", "#,##0.00")
    line("cpr", "Analyst cost per reviewed case", "BDT", "[cpm]*[mins]", "#,##0.00")
    line("ana", "Analyst time cost", "BDT", "[reviews]*[cpr]", "#,##0", "reviews x cost per review")
    line("fric", "Customer-friction cost", "BDT", "[fp]*([aband]*[rev]+[churn]*[cval])", "#,##0", "false alarms x (abandonment x lost revenue + churn x customer margin)")
    line("run", "Run cost", "BDT", "[opex]", "#,##0")
    line("fte", "Analysts needed", "FTE", "[reviews]*[mins]/([hrs]*60*[util])", "#,##0.0", "capacity check: can the team really review this many cases?")

    section("D. Result")
    line("net", "NET MONTHLY BENEFIT", "BDT", "[prev]-[fa]-[ana]-[fric]-[run]", "#,##0;(#,##0);-", "prevented loss - false-alarm - analyst - friction - run", key_row=True)
    line("pay", "PAYBACK", "months", 'IF([net]>0,[capex]/[net],"never")', "#,##0.0", "one-time cost / net monthly benefit; 'never' if net <= 0", key_row=True)
    line("roi", "ROI over horizon", "fraction", "([net]*[hor]-[capex])/[capex]", "0%;(0%);-", "(net x horizon - one-time) / one-time", key_row=True)
    line("cum", "Cumulative net over horizon, after one-time cost", "BDT", "[net]*[hor]-[capex]", "#,##0;(#,##0);-")

    section("E. Break-even checks")
    line("cfa", "All-in cost of ONE false alarm", "BDT", "[notif]+[scr]*[scc]+[aband]*[rev]+[churn]*[cval]+[share]*[cpr]", "#,##0.00", "compare with the flat 20 BDT used in Step 2 (config/eval_costs.json)")
    line("gtp", "Net value of one flagged scam", "BDT", "[L]*[e]-[share]*[cpr]", "#,##0.00", "avoided loss minus its own review cost")
    line("befar", "Break-even false-alarm rate", "fraction", 'IF(([N]*[s]*[r]*[gtp]-[run])>0,([N]*[s]*[r]*[gtp]-[run])/([N]*(1-[s])*[cfa]),"none")', "0.000%", "highest false-alarm rate at which net = 0; 'none' = loses money even with zero false alarms")
    line("bes", "Break-even scam rate", "fraction", 'IF(([N]*([r]*[gtp]+[far]*[cfa]))>0,([N]*[far]*[cfa]+[run])/([N]*([r]*[gtp]+[far]*[cfa])),"none")', "0.000%", "lowest scam rate at which net = 0")
    line("pcl", "Share of total scam loss prevented", "fraction", "IF([gross]=0,0,[prev]/[gross])", "0.0%")
    wm.freeze_panes = "C5"

    # ---------------- Sensitivity ----------------
    wsn = wb.create_sheet("Sensitivity")
    setw(wsn, [30, 16, 16, 16, 16, 16, 16, 16])
    wsn["A1"].value, wsn["A1"].font = "Sensitivity of monthly net benefit (BDT), Medium scenario", f_title
    wsn["A2"].value, wsn["A2"].font = "Axis values (blue) are editable. Every other input comes from ScamShield_Model column D (Medium).", f_warn
    D = lambda k: f"ScamShield_Model!$D${mrow[k]}"
    put(wsn, "A4", "Grid 1: recall (across) x false-alarm rate (down)", f_bold, border=False)
    put(wsn, "A5", "false-alarm rate \\ recall", f_bold)
    for j, rc in enumerate(RECALL_AXIS):
        put(wsn, f"{L_(2+j)}5", rc, f_in, "0.0%", fill_in)
    for i, fa in enumerate(FAR_AXIS):
        rr = 6 + i
        put(wsn, f"A{rr}", fa, f_in, "0.00%", fill_in)
        for j in range(len(RECALL_AXIS)):
            cl = L_(2 + j)
            fml = (f"={D('N')}*{D('s')}*{cl}$5*({D('L')}*{D('e')}-{D('share')}*{D('cpr')})"
                   f"-{D('N')}*(1-{D('s')})*$A{rr}*{D('cfa')}-{D('run')}")
            put(wsn, f"{cl}{rr}", fml, f_norm, "#,##0;(#,##0);-")
    g2 = 6 + len(FAR_AXIS) + 2
    put(wsn, f"A{g2}", "Grid 2: scam rate (across) x average loss (down)", f_bold, border=False)
    put(wsn, f"A{g2+1}", "average loss \\ scam rate", f_bold)
    for j, sr in enumerate(RATE_AXIS):
        put(wsn, f"{L_(2+j)}{g2+1}", sr, f_in, "0.000%", fill_in)
    for i, lo in enumerate(LOSS_AXIS):
        rr = g2 + 2 + i
        put(wsn, f"A{rr}", lo, f_in, "#,##0", fill_in)
        for j in range(len(RATE_AXIS)):
            cl = L_(2 + j)
            fml = (f"={D('N')}*{cl}${g2+1}*{D('r')}*($A{rr}*{D('e')}-{D('share')}*{D('cpr')})"
                   f"-{D('N')}*(1-{cl}${g2+1})*{D('far')}*{D('cfa')}-{D('run')}")
            put(wsn, f"{cl}{rr}", fml, f_norm, "#,##0;(#,##0);-")
    put(wsn, f"A{g2+2+len(LOSS_AXIS)+1}", "Positive = ScamShield earns more than it costs per month. One-at-a-time +/-25% sensitivity of every input is in results/step8_sensitivity_tornado.csv.", f_norm, border=False)

    # ---------------- AgentCopilot_KPIs ----------------
    wk = wb.create_sheet("AgentCopilot_KPIs")
    setw(wk, [30, 26, 10, 54, 54, 40, 40, 14, 14, 14, 14, 14, 14])
    wk["A1"].value, wk["A1"].font = "Agent Copilot: measurable KPIs", f_title
    wk["A2"].value, wk["A2"].font = ("Baseline / pilot / target columns hold EXAMPLE values (placeholders; forecast-error row = our synthetic Step 1 result). "
                                    "Replace them with measured values. Design: randomise agents (stratified by area and size) into Copilot on / off for 8 weeks."), f_warn
    head(wk, 4, ["KPI", "Unit", "Better", "Definition", "Formula", "Data needed", "Guardrail", "Baseline (EXAMPLE)", "Pilot (EXAMPLE)", "Target change (ASSUMPTION)", "Change", "Target met?", "Status"])
    for i, k in enumerate(KPIS):
        rr = 5 + i
        put(wk, f"A{rr}", k["kpi"], f_bold, align=wrap)
        put(wk, f"B{rr}", k["unit"], align=wrap); put(wk, f"C{rr}", k["better"])
        put(wk, f"D{rr}", k["definition"], align=wrap); put(wk, f"E{rr}", k["formula"], align=wrap)
        put(wk, f"F{rr}", k["data_needed"], align=wrap); put(wk, f"G{rr}", k["guardrail"], align=wrap)
        is_share = k["kpi"] == "Warning usefulness"
        put(wk, f"H{rr}", k["baseline_example"], f_in, "0%" if is_share else "#,##0.0", fill_in)
        put(wk, f"I{rr}", k["pilot_example"], f_in, "0%" if is_share else "#,##0.0", fill_in)
        put(wk, f"J{rr}", k["target_example"], f_in, "+0%;-0%", fill_in)
        if is_share:
            put(wk, f"K{rr}", f"=I{rr}-H{rr}", f_norm, "+0%;-0%;0%")               # percentage-point change
            ok = f'=IF(C{rr}="higher",K{rr}>=J{rr},K{rr}<=J{rr})'
        else:
            put(wk, f"K{rr}", f"=IF(H{rr}=0,0,I{rr}/H{rr}-1)", f_norm, "+0.0%;-0.0%;0%")   # relative change
            ok = f'=IF(C{rr}="higher",K{rr}>=J{rr},K{rr}<=J{rr})'
        put(wk, f"L{rr}", ok, f_norm)
        put(wk, f"M{rr}", "EXAMPLE", f_warn)
        wk.row_dimensions[rr].height = 92
    nr = 5 + len(KPIS) + 1
    put(wk, f"A{nr}", "Notes", f_bold, border=False)
    notes = ["'Change' is relative (pilot / baseline - 1) except 'Warning usefulness', which is the percentage-point difference.",
             "Read stock-out hours, completed transactions and idle cash TOGETHER: lower idle cash with more stock-outs is a failure.",
             "The synthetic dataset has no cash balances, so these KPIs can only be measured in a real pilot with hourly balance snapshots and failure-reason codes.",
             "Sources to cite for context and baselines: S14 (agent-network surveys), S1 (Bangladesh Bank MFS statistics), S15-S16 (experiment design), S11 (cost of funds)."]
    for i, t in enumerate(notes):
        put(wk, f"A{nr+1+i}", t, f_norm, border=False)

    # ---------------- CostAdvisor_ABPilot ----------------
    wc = wb.create_sheet("CostAdvisor_ABPilot")
    setw(wc, [58, 22, 16, 24, 70, 50])
    wc["A1"].value, wc["A1"].font = "Cost Advisor: A/B pilot to measure REAL fee savings", f_title
    wc["A2"].value, wc["A2"].font = "No migration rate is assumed anywhere. Migration and savings are measured outputs of the pilot (section 3).", f_warn
    put(wc, "A4", "1. Design", f_bold, border=False)
    head(wc, 5, ["Element", "Design"])
    wc.merge_cells("B5:E5")
    for i, (el, ds) in enumerate(AB_DESIGN):
        rr = 6 + i
        put(wc, f"A{rr}", el, f_bold, align=wrap)
        wc.merge_cells(f"B{rr}:E{rr}")
        put(wc, f"B{rr}", ds, align=wrap)
        wc.row_dimensions[rr].height = 46
    s2 = 6 + len(AB_DESIGN) + 1
    put(wc, f"A{s2}", "2. Sample size (edit blue cells)", f_bold, border=False)
    head(wc, s2 + 1, ["Input", "Value", "Unit", "Status", "Basis", "Public sources to cite"])
    brow = {}
    for i, b in enumerate(cfg["ab_pilot"]):
        rr = s2 + 2 + i
        brow[b["id"]] = rr
        fmt = "0.0%" if b["unit"].startswith("fraction") and b["id"] not in ("B9", "B10", "B11") else ("0.0%" if b["unit"].startswith("fraction") else "#,##0.0")
        if b["unit"] in ("customers", "BDT", "weeks") and b["id"] != "B3" and b["id"] != "B4":
            fmt = "#,##0"
        put(wc, f"A{rr}", f"{b['id']}  {b['name']}", align=wrap)
        put(wc, f"B{rr}", b["value"], f_in, fmt, fill_in)
        put(wc, f"C{rr}", b["unit"], align=wrap)
        put(wc, f"D{rr}", b["kind"], f_bold if b["kind"] != "ASSUMPTION" else f_norm, align=wrap)
        put(wc, f"E{rr}", b["basis"], align=wrap)
        put(wc, f"F{rr}", source_text(cfg, b["sources"]), align=wrap)
        wc.row_dimensions[rr].height = 52
    r0 = s2 + 2 + len(cfg["ab_pilot"])
    B = lambda i: f"$B${brow[i]}"
    put(wc, f"A{r0}", "z (alpha/2) + z (power)", f_norm)
    put(wc, f"B{r0}", f"=NORMSINV(1-{B('B1')}/2)+NORMSINV({B('B2')})", f_norm, "0.000")
    put(wc, f"A{r0+1}", "Customers needed per arm", f_bold)
    put(wc, f"B{r0+1}", f"=ROUNDUP(2*B{r0}^2*{B('B3')}^2*(1-{B('B5')}^2)/{B('B4')}^2,0)", f_bold, "#,##0", fill_key)
    put(wc, f"E{r0+1}", "n = 2 (z_a/2 + z_b)^2 sigma^2 (1 - rho^2) / MDE^2   (normal approximation)", align=wrap)
    put(wc, f"A{r0+2}", "Customers to enrol (both arms, after attrition)", f_bold)
    put(wc, f"B{r0+2}", f"=ROUNDUP(2*B{r0+1}/(1-{B('B6')}),0)", f_bold, "#,##0", fill_key)
    put(wc, f"A{r0+3}", "Enough eligible customers?", f_bold)
    put(wc, f"B{r0+3}", f'=IF({B("B7")}>=B{r0+2},"yes","no: raise MDE, extend duration or widen eligibility")', f_norm)
    s3 = r0 + 5
    put(wc, f"A{s3}", "3. Analysis template (enter pilot results; values below are EXAMPLES)", f_bold, border=False)
    head(wc, s3 + 1, ["Quantity", "Value", "Unit", "Status", "How to get it"])
    ex = [
        ("n_t", "Treatment customers analysed", 1100, "customers", "EXAMPLE", "count after attrition"),
        ("n_c", "Control customers analysed", 1090, "customers", "EXAMPLE", "count after attrition"),
        ("m_t", "Mean fees paid per customer-month, treatment", 27.0, "BDT", "EXAMPLE", "fees on cash-outs, pilot period (CUPED-adjusted if used)"),
        ("m_c", "Mean fees paid per customer-month, control", 29.0, "BDT", "EXAMPLE", "same definition"),
        ("sd_t", "Std. deviation, treatment", 19.0, "BDT", "EXAMPLE", "observed"),
        ("sd_c", "Std. deviation, control", 19.0, "BDT", "EXAMPLE", "observed"),
        ("val", "Eligible small cash-out value per customer-month (control)", 2100.0, "BDT", "EXAMPLE", "observed in control arm; cash-outs below the small limit"),
        ("mig", "Observed migration (share of eligible small cash-out value moved to digital)", 0.05, "fraction", "EXAMPLE", "MEASURED in pilot = (control small cash-out value - treatment small cash-out value) / control; NOT assumed"),
    ]
    arw = {}
    for i, (k, lab, v, unit, st, how) in enumerate(ex):
        rr = s3 + 2 + i
        arw[k] = rr
        put(wc, f"A{rr}", lab, align=wrap)
        put(wc, f"B{rr}", v, f_in, "0.0%" if unit == "fraction" else "#,##0.0", fill_in)
        put(wc, f"C{rr}", unit); put(wc, f"D{rr}", st, f_warn); put(wc, f"E{rr}", how, align=wrap)
        wc.row_dimensions[rr].height = 30
    q = s3 + 2 + len(ex)
    X = lambda k: f"$B${arw[k]}"
    calc = [
        ("diff", "Difference in mean fees (treatment - control)", f"={X('m_t')}-{X('m_c')}", "#,##0.00;(#,##0.00)", "negative = customers pay less"),
        ("se", "Standard error", f"=SQRT({X('sd_t')}^2/{X('n_t')}+{X('sd_c')}^2/{X('n_c')})", "0.000", ""),
        ("save", "Customer saving per customer-month (BDT)", f"=-B{q}", "#,##0.00;(#,##0.00)", "equals minus the difference"),
        ("lo", "95% CI lower bound of saving", f"=B{q+2}-1.96*B{q+1}", "#,##0.00;(#,##0.00)", "use this lower bound for scale-up forecasts"),
        ("hi", "95% CI upper bound of saving", f"=B{q+2}+1.96*B{q+1}", "#,##0.00;(#,##0.00)", ""),
        ("sig", "Statistically significant saving?", f'=IF(B{q+3}>0,"yes","no")', None, "lower bound above zero"),
        ("mde", "Lower bound above the minimum worth acting on?", f'=IF(B{q+3}>={B("B4")},"yes","no")', None, "compares with MDE (B4)"),
        ("xc", "Cross-check: saving implied by observed migration", f"={X('val')}*{X('mig')}*({B('B9')}-{B('B10')})", "#,##0.00", "eligible value x observed migration x (cash-out fee rate - digital fee rate); should be close to the saving above"),
        ("rev", "upay net revenue change per customer-month", f"=({X('val')}*{X('mig')})*({B('B10')}+{B('B11')}-{B('B9')})", "#,##0.00;(#,##0.00)", "migrated value x (digital fee + merchant discount - cash-out fee); negative = revenue transfer to customers"),
    ]
    for i, (k, lab, fml, fmt, how) in enumerate(calc):
        rr = q + i
        put(wc, f"A{rr}", lab, f_bold if k in ("save", "lo") else f_norm, align=wrap)
        put(wc, f"B{rr}", fml, f_bold if k in ("save", "lo") else f_norm, fmt, fill_key if k in ("save", "lo") else None)
        put(wc, f"E{rr}", how, align=wrap)
    put(wc, f"A{q+len(calc)+1}", "Do not enter a migration rate before the pilot. Scale-up value = eligible customers x lower-bound saving x 12, never a fixed rate.", f_warn, border=False)

    # ---------------- Sources ----------------
    wsrc = wb.create_sheet("Sources")
    setw(wsrc, [6, 62, 44, 70, 22, 56])
    wsrc["A1"].value, wsrc["A1"].font = "Public sources to cite for each input", f_title
    wsrc["A2"].value, wsrc["A2"].font = "Only S1 was seen live (2026-10-07). For every other source open it, copy the exact figure, edition and date, and cite that. Do not cite a source for a number you did not read in it.", f_warn
    head(wsrc, 4, ["ID", "Source", "Where", "What to take from it", "Used for inputs", "Note"])
    used = {}
    for a in cfg["assumptions"] + cfg["ab_pilot"]:
        for sid in a["sources"]:
            used.setdefault(sid, []).append(a["id"])
    for i, (sid, s) in enumerate(cfg["sources"].items()):
        rr = 5 + i
        put(wsrc, f"A{rr}", sid, f_bold)
        put(wsrc, f"B{rr}", s["title"], align=wrap); put(wsrc, f"C{rr}", s["url"], align=wrap)
        put(wsrc, f"D{rr}", s["use"], align=wrap); put(wsrc, f"E{rr}", ", ".join(used.get(sid, [])), align=wrap)
        put(wsrc, f"F{rr}", source_note(sid), align=wrap)
        wsrc.row_dimensions[rr].height = 48

    for sheet in wb.worksheets:
        sheet.sheet_view.showGridLines = False
    wb.save(path)
    return {"model_rows": mrow, "assumption_rows": arow, "ab_rows": brow, "ab_calc_start": q, "ab_r0": r0}


# ----------------------------------------------------------------------------------------------
def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--config", default=str(CONFIG))
    ap.add_argument("--out-dir", default=str(ROOT / "results"))
    ap.add_argument("--no-xlsx", action="store_true", help="skip the workbook (CSV and Markdown only)")
    a = ap.parse_args(argv)
    cfg = load_config(Path(a.config))
    out = Path(a.out_dir)
    out.mkdir(parents=True, exist_ok=True)
    res = export_csvs(cfg, out)
    write_markdown(cfg, res, out)
    if not a.no_xlsx:
        build_xlsx(cfg, out / "step8_impact_model.xlsx")
    for s in SCEN:
        r = res[s]
        pb = "never" if r["payback_months"] is None else f"{r['payback_months']:.1f} months"
        print(f"{s:7s} prevented {r['prevented_loss']/1e6:8.2f} M | net/month {r['net_benefit']/1e6:8.2f} M | payback {pb} | "
              f"false alarms {r['fp']:,.0f} | analysts {r['analyst_fte']:.1f}")
    return res


if __name__ == "__main__":
    sys.exit(0 if main() is not None else 1)
